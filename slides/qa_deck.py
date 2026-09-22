"""
Geometry QA without a renderer.

Estimates wrapped text height from character-width tables for Roboto/Raleway at
the given point size, then flags boxes whose text cannot fit, shapes outside the
safe margin, and overlapping text boxes.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

DECK = sys.argv[1] if len(sys.argv) > 1 else r"C:\E\VIT\VIT 7th Sem\Deep Learning Project\Review-2.pptx"
EMU_IN = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5

# Average glyph width as a fraction of font size, measured for these families.
AVG_CHAR = {"Raleway": 0.50, "Roboto": 0.505, "Consolas": 0.55}


def est_height(tf, box_w_in, default_font="Roboto"):
    """Estimated rendered height in inches."""
    total = 0.0
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            total += 0.14
            continue
        size = max((r.font.size.pt if r.font.size else 11) for r in runs)
        font = next((r.font.name for r in runs if r.font.name), default_font)
        bold = any(r.font.bold for r in runs)
        txt = "".join(r.text for r in runs)
        cw = AVG_CHAR.get(font, 0.505) * (1.05 if bold else 1.0) * size / 72.0
        per_line = max(1, int(box_w_in / cw)) if cw else 999
        lines = max(1, -(-len(txt) // per_line))
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.18
        total += lines * size * ls / 72.0
    return total


prs = Presentation(DECK)
issues = []
for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        x, y = sh.left / EMU_IN, sh.top / EMU_IN
        w, h = sh.width / EMU_IN, sh.height / EMU_IN

        if x < MARGIN - 0.03 or y < MARGIN - 0.25 or x + w > SLIDE_W - MARGIN + 0.03 or y + h > SLIDE_H - MARGIN + 0.03:
            issues.append(f"S{idx:>2} MARGIN   '{sh.name[:22]}' at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f} "
                          f"-> right={x+w:.2f} bottom={y+h:.2f}")

        if sh.has_text_frame and sh.text_frame.text.strip():
            need = est_height(sh.text_frame, w - 0.08)
            if need > h + 0.06:
                issues.append(f"S{idx:>2} OVERFLOW '{sh.name[:22]}' needs {need:.2f}\" has {h:.2f}\" "
                              f": {sh.text_frame.text.strip()[:58]!r}")
            boxes.append((x, y, w, max(h, need), sh.name, sh.text_frame.text.strip()[:30]))

    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.06 and oy > 0.06:
                issues.append(f"S{idx:>2} OVERLAP  '{a[4][:18]}'({a[5][:18]!r}) x "
                              f"'{b[4][:18]}'({b[5][:18]!r}) by {ox:.2f}x{oy:.2f}\"")

print(f"{len(prs.slides._sldIdLst)} slides, {len(issues)} issues\n")
for i in issues:
    print(" ", i)
